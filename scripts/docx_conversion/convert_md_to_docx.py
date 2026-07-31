#!/usr/bin/env python3
"""
convert_md_to_docx.py — Markdown → .docx with real Word table design.

Pipeline:
  [1/2] pandoc converts .md → .docx (preserves EVERYTHING: headings, tables,
        code blocks, blockquotes, lists, emoji, inline formatting — no detail lost).
  [2/2] python-docx post-processes every table to give it Word table design:
        thin gray borders on all cells, bold + subtly-shaded header row,
        autofit columns to content.

Usage:
    py scripts/docx_conversion/convert_md_to_docx.py <input.md> [<output.docx>]

If output path is omitted, it uses <input>.docx in the same directory.
The script is content-preserving: nothing from the .md is dropped. The only
post-processing is visual styling of tables (borders, header band).
"""

import sys
import subprocess
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ─── Step 1: pandoc base conversion ────────────────────────────────────────

def run_pandoc(input_md: Path, output_docx: Path) -> None:
    """Run pandoc to convert .md → .docx (preserves everything)."""
    cmd = [
        'pandoc',
        str(input_md),
        '-o', str(output_docx),
        '--wrap=preserve',  # don't rewrap lines
        '--toc',             # table of contents
        '--toc-depth=3',     # up to ###
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(f"pandoc failed (code {result.returncode}):\n{result.stderr}")
    print(f"[1/2] pandoc: {input_md.name} → {output_docx.name}")


# ─── Step 2: python-docx table polish ──────────────────────────────────────

def _set_cell_borders(cell, color: str = '6F6F6F', sz: int = 4) -> None:
    """Apply thin gray borders to all four sides of a cell (sz in 1/8 pt)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    # Remove pre-existing tcBorders before creating a fresh one.
    existing = tcPr.find(qn('w:tcBorders'))
    if existing is not None:
        tcPr.remove(existing)
    tcBorders = OxmlElement('w:tcBorders')
    tcPr.append(tcBorders)
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), str(sz))
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), color)
        tcBorders.append(b)


def _set_cell_shading(cell, fill: str = 'E7E6E6') -> None:
    """Apply a subtle gray background fill to a cell (header row band)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    existing = tcPr.find(qn('w:shd'))
    if existing is not None:
        tcPr.remove(existing)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def style_table(table) -> None:
    """Apply Word table design: borders, bold+s shaded header, autofit."""
    # Borders on every cell.
    for row in table.rows:
        for cell in row.cells:
            _set_cell_borders(cell)
    # Header row (first): bold + shaded.
    if table.rows:
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
            _set_cell_shading(cell)
    # Autofit columns to content.
    table.autofit = True


def polish_tables(docx_path: Path) -> int:
    """Step 2: Walk every table in the .docx and apply style_table()."""
    doc = Document(str(docx_path))
    count = 0
    for table in doc.tables:
        style_table(table)
        count += 1
    doc.save(str(docx_path))
    print(f"[2/2] python-docx: polished {count} tables (borders + header band + autofit)")
    return count


# ─── Entry point ────────────────────────────────────────────────────────────

def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    input_md = Path(sys.argv[1])
    output_docx = Path(sys.argv[2]) if len(sys.argv) >= 3 else input_md.with_suffix('.docx')
    if not input_md.is_file():
        print(f"ERROR: input not found: {input_md}", file=sys.stderr)
        sys.exit(1)
    if not output_docx.suffix.lower() == '.docx':
        output_docx = output_docx.with_suffix('.docx')

    run_pandoc(input_md, output_docx)
    table_count = polish_tables(output_docx)

    # Brief report.
    size_kb = output_docx.stat().st_size / 1024
    print(f"\n✓ Done: {input_md.name} → {output_docx} ({size_kb:.1f} KB, {table_count} tables styled)")


if __name__ == '__main__':
    main()
